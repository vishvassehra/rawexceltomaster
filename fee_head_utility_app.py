import streamlit as st
from io import BytesIO
from copy import copy
import openpyxl

st.set_page_config(page_title="Fee Head Wise Excel Utility", page_icon="📊", layout="centered")

st.title("📊 Fee Head Wise Excel Utility")
st.write("Upload your fee collection report and convert each fee head with an amount into a separate row.")

st.info(
    "Output: Receipt No. | Adm No. | Name | Class | Receipt Date | "
    "Installment | Receipt Mode | Fee Head | Amount\n\n"
    "Paid Amount and payableAmount are not included."
)

uploaded_file = st.file_uploader(
    "Upload Excel report",
    type=["xlsx", "xlsm"]
)

if uploaded_file:
    st.success(f"Selected: {uploaded_file.name}")

    if st.button("⚙️ Process & Create Fee Head Wise File", type="primary", use_container_width=True):
        try:
            keep_vba = uploaded_file.name.lower().endswith(".xlsm")
            wb = openpyxl.load_workbook(BytesIO(uploaded_file.getvalue()), data_only=True)

            output_wb = openpyxl.Workbook()
            output_ws = output_wb.active
            output_ws.title = "Fee Head Wise"

            headers = [
                "Receipt No.", "Adm No.", "Name", "Class", "Receipt Date",
                "Installment", "Receipt Mode", "Fee Head", "Amount"
            ]
            output_ws.append(headers)

            total_records = 0

            for ws in wb.worksheets:
                # Find the report header row by looking for Receipt No.
                header_row = None
                for row in ws.iter_rows():
                    for cell in row:
                        if str(cell.value).strip().lower() == "receipt no.":
                            header_row = cell.row
                            break
                    if header_row:
                        break

                if not header_row:
                    continue

                # Identify important columns from the report header.
                header_map = {}
                for c in range(1, ws.max_column + 1):
                    value = ws.cell(header_row, c).value
                    if value is not None:
                        header_map[str(value).strip().lower()] = c

                receipt_col = header_map.get("receipt no.", 1)
                adm_col = header_map.get("adm no.", 2)
                name_col = header_map.get("name", 3)
                class_col = header_map.get("class", 4)
                date_col = header_map.get("receipt date", 5)
                installment_col = header_map.get("installment", 6)
                mode_col = header_map.get("receipt mode", 7)

                # Fee heads are the columns after Receipt Mode and before Total.
                total_col = header_map.get("total", ws.max_column + 1)
                fee_start = mode_col + 1
                fee_end = total_col - 1

                current_student = [None] * 5

                for r in range(header_row + 1, ws.max_row + 1):
                    row_has_data = any(
                        ws.cell(r, c).value is not None
                        and str(ws.cell(r, c).value).strip() != ""
                        for c in range(1, ws.max_column + 1)
                    )
                    if not row_has_data:
                        continue

                    # Carry forward student details from above whenever source cells are blank.
                    source_cols = [receipt_col, adm_col, name_col, class_col, date_col]
                    for i, c in enumerate(source_cols):
                        value = ws.cell(r, c).value
                        if value is not None and str(value).strip() != "":
                            current_student[i] = value

                    installment = ws.cell(r, installment_col).value
                    receipt_mode = ws.cell(r, mode_col).value

                    # One output row for every fee head that has a non-zero amount.
                    for c in range(fee_start, fee_end + 1):
                        head = ws.cell(header_row, c).value
                        amount = ws.cell(r, c).value

                        if head is None or str(head).strip() == "":
                            continue
                        if amount is None or (isinstance(amount, str) and amount.strip() == ""):
                            continue

                        # "Amount hai" means zero values are not exported.
                        try:
                            if float(amount) == 0:
                                continue
                        except (ValueError, TypeError):
                            pass

                        output_ws.append(
                            current_student + [
                                installment,
                                receipt_mode,
                                str(head).strip(),
                                amount
                            ]
                        )
                        total_records += 1

            # Formatting
            for cell in output_ws[1]:
                cell.font = copy(output_ws[1][1].font)
                cell.alignment = copy(output_ws[1][1].alignment)

            widths = [14, 14, 22, 16, 18, 22, 28, 28, 14]
            for i, width in enumerate(widths, 1):
                output_ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

            output_ws.freeze_panes = "A2"
            output_ws.auto_filter.ref = output_ws.dimensions

            if total_records == 0:
                st.warning("No fee-head amounts were found.")
            else:
                result = BytesIO()
                output_wb.save(result)
                result.seek(0)

                base = uploaded_file.name.rsplit(".", 1)[0]
                output_name = f"{base}_Fee_Head_Wise.xlsx"

                st.success(f"Done! Created {total_records:,} fee-head rows.")

                st.download_button(
                    "⬇️ Download Fee Head Wise Excel",
                    data=result.getvalue(),
                    file_name=output_name,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )

        except Exception as e:
            st.error("The file could not be processed.")
            st.exception(e)

st.divider()
st.caption("Each fee head with a non-zero amount becomes one row. Existing student details are carried forward from the previous row.")
